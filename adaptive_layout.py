"""Adaptive, print-safe attendance workbook layout for GitHub Pages / Pyodide.

All dimensions are designed in physical millimetres before translation to Excel
column units. Week and exam groups are never cut in half across sheets. Student
rows paginate vertically with the same student cohort repeated on each sheet.
"""
from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.page import PageMargins

import yoklama_olustur as core

# A4 landscape: 297 × 210 mm; horizontal paper margins 4 mm each.
PAGE_WIDTH_MM = 297.0
SIDE_MARGIN_MM = 4.0
PRINTABLE_MM = PAGE_WIDTH_MM - 2 * SIDE_MARGIN_MM
# Initial column widths are not fixed as a group: name width is data-dependent.
SERIAL_WIDTH_MM = 7.5
STUDENT_NUMBER_WIDTH_MM = 21.0
NAME_MIN_MM = 28.0
NAME_MAX_DEFAULT_MM = 58.0              # preserves the previous name-column ceiling
FIXED_MM = (SERIAL_WIDTH_MM, STUDENT_NUMBER_WIDTH_MM, NAME_MAX_DEFAULT_MM)  # conservative preview
MARK_MIN_MM = 15.0                     # user-selectable default, physical handwriting minimum
MARK_MAX_MM = 18.0                     # may increase if user chooses minimum > 18 mm
MAX_COLUMNS_WIDTH_MM = PRINTABLE_MM - sum(FIXED_MM)
FONT_DEFAULT_PT = 8                    # previous student-name font: 10 pt
FONT_FAMILIES = ('Aptos', 'Arial', 'Calibri', 'Times New Roman')
HEADER_ROWS = 6
STUDENTS_PER_PAGE = 17
STUDENT_HEIGHT_PT = 23.5              # 8.29 mm
MIN_STUDENT_HEIGHT_PT = 23.5
MAX_STUDENT_HEIGHT_PT = 27.0

NAVY = "222239"
BURGUNDY = "9B0D22"
PALE = "F7E6E9"
GRAY = "F0F1F3"
EXAM_GRAY = "D3D4DA"
WHITE = "FFFFFF"
TEXT = "24202A"
LIGHT_LINE = "A6A6AD"


@dataclass(frozen=True)
class Group:
    week: int
    kind: str             # week / exam
    label: str
    span: int             # signature columns, or 1 for an exam


def validate_settings(weeks, signatures, midterm, final):
    try:
        weeks, signatures, midterm, final = (int(x) for x in (weeks, signatures, midterm, final))
    except (TypeError, ValueError) as exc:
        raise ValueError("Hafta, imza, vize ve final seçimleri tam sayı olmalıdır") from exc
    if not 10 <= weeks <= 20:
        raise ValueError("Hafta sayısı 10 ile 20 arasında olmalıdır")
    if not 1 <= signatures <= 5:
        raise ValueError("İmza sayısı 1 ile 5 arasında olmalıdır")
    if not 1 <= midterm < final <= weeks:
        raise ValueError("Sınav haftaları 1 ≤ Vize < Final ≤ Hafta sayısı şartını sağlamalıdır")
    return weeks, signatures, midterm, final


def groups_for_weeks(weeks, signatures, midterm, final):
    groups = []
    for week in range(1, weeks + 1):
        if week == midterm:
            groups.append(Group(week, 'exam', f'VİZE SINAVI · {week}. HAFTA', 1))
        elif week == final:
            groups.append(Group(week, 'exam', f'FİNAL SINAVI · {week}. HAFTA', 1))
        else:
            groups.append(Group(week, 'week', f'{week}. HAFTA', signatures))
    return groups


def validate_presentation(min_mark_mm=MARK_MIN_MM, font_size_pt=FONT_DEFAULT_PT,
                          font_family='Aptos', name_max_mm=NAME_MAX_DEFAULT_MM):
    """Validate display settings once; never trust values coming from the browser."""
    try:
        min_mark_mm = float(min_mark_mm)
        font_size_pt = int(font_size_pt)
        name_max_mm = float(name_max_mm)
    except (TypeError, ValueError) as exc:
        raise ValueError('Yazı puntosu ve genişlikler sayısal olmalıdır') from exc
    if not (math.isfinite(min_mark_mm) and 8 <= min_mark_mm <= 20):
        raise ValueError('Asgari imza genişliği 8–20 mm olmalıdır')
    if not 6 <= font_size_pt <= 14:
        raise ValueError('Yazı puntosu 6–14 arasında olmalıdır')
    if not (math.isfinite(name_max_mm) and 30 <= name_max_mm <= 80):
        raise ValueError('Ad Soyad üst sınırı 30–80 mm olmalıdır')
    if font_family not in FONT_FAMILIES:
        raise ValueError('Desteklenmeyen yazı tipi')
    return min_mark_mm, font_size_pt, font_family, name_max_mm


def raw_name_width_mm(name, font_size_pt=FONT_DEFAULT_PT, font_family='Aptos'):
    """Approximate bold text length without OS-dependent font rendering."""
    def glyph_em(ch):
        if ch.isspace(): return .34
        if ch in 'Iİıl|!.,;:\'': return .31
        if ch in 'MWŞĞÖÜÇQ@%': return .77
        if ch.isdigit(): return .56
        if unicodedata.east_asian_width(ch) in ('W', 'F'): return 1.0
        if ch.isupper(): return .61
        return .52
    family_factor = 0.92 if font_family == 'Aptos' else (1.0 if font_family == 'Times New Roman' else .97)
    return sum(glyph_em(ch) for ch in str(name or '')) * font_size_pt * (25.4 / 72) * family_factor


def estimated_name_width_mm(students, font_size_pt=FONT_DEFAULT_PT, maximum_mm=NAME_MAX_DEFAULT_MM,
                            font_family='Aptos'):
    """The longest student's name determines the entire Name-column width."""
    widest = max(raw_name_width_mm(name, font_size_pt, font_family) for _, name in students)
    return round(max(NAME_MIN_MM, min(maximum_mm, widest + 4.0)), 2)


def fitted_name_font_pt(name, width_mm, requested_pt, font_family='Aptos'):
    """Only downsize an exceptionally long name so it fits in at most two lines."""
    max_one_line = max(1, width_mm - 3.5)
    if raw_name_width_mm(name, requested_pt, font_family) <= max_one_line:
        return requested_pt
    # A 23.5 pt row contains up to two lines at approximately 9 pt.
    for candidate in range(min(requested_pt, 9), 5, -1):
        if raw_name_width_mm(name, candidate, font_family) <= max_one_line * 1.95:
            return candidate
    return 6


def page_plan(weeks, signatures, midterm, final, min_mark_mm=MARK_MIN_MM, fixed_widths_mm=FIXED_MM):
    """Group full weeks into A4 landscape sheets with a physical minimum cell width."""
    weeks, signatures, midterm, final = validate_settings(weeks, signatures, midterm, final)
    if not math.isfinite(min_mark_mm) or not 8 <= min_mark_mm <= 20:
        raise ValueError('Asgari imza genişliği 8–20 mm olmalıdır')
    available = PRINTABLE_MM - sum(fixed_widths_mm)
    if available <= 0:
        raise ValueError('Sütunlara ayrılan alan A4 genişliğini aşıyor')
    all_groups = groups_for_weeks(weeks, signatures, midterm, final)
    # Dynamic programming: minimize A4 sections first, then balance their
    # signature-column counts. A greedy fill strands tiny, unreadable last pages.
    n = len(all_groups)
    dp = [None] * (n + 1)
    dp[n] = (0, 0, [])
    for start in range(n - 1, -1, -1):
        used = 0
        best = None
        for end in range(start + 1, n + 1):
            used += all_groups[end - 1].span
            if used * min_mark_mm > available + 1e-9:
                break
            remaining = dp[end]
            if remaining is None:
                continue
            candidate = (remaining[0] + 1, remaining[1] + used ** 2,
                         [all_groups[start:end]] + remaining[2])
            if best is None or candidate[:2] < best[:2]:
                best = candidate
        dp[start] = best
    if dp[0] is None:
        raise ValueError('Geçerli bir A4 sayfa planı oluşturulamadı')
    return dp[0][2]

def mm_to_excel_chars(mm):
    """Calibri/Aptos Narrow column conversion at nominal 96 dpi (7px/char)."""
    return round((mm * 96 / 25.4 - 5) / 7, 2)


def mark_width_mm(groups, min_mark_mm=MARK_MIN_MM, fixed_widths_mm=FIXED_MM):
    count = sum(g.span for g in groups)
    if not count:
        raise ValueError('Çıktı sayfası boş olamaz')
    available = PRINTABLE_MM - sum(fixed_widths_mm)
    value = min(max(MARK_MAX_MM, min_mark_mm), available / count)
    if value + 1e-7 < min_mark_mm:
        raise ValueError('İmza hücresi için yeterli genişlik yok')
    return round(value, 4)


def _print_setup(ws, last_column, last_row):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=SIDE_MARGIN_MM/25.4, right=SIDE_MARGIN_MM/25.4,
                                   top=5/25.4, bottom=5/25.4, header=0, footer=0)
    ws.print_title_rows = '1:6'
    ws.print_area = f'A1:{get_column_letter(last_column)}{last_row}'
    for after in range(HEADER_ROWS + STUDENTS_PER_PAGE, last_row, STUDENTS_PER_PAGE):
        ws.row_breaks.append(Break(id=after))
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 75
    ws.freeze_panes = 'D7'


def _safe_text(text):
    # Avoid Excel treating student-supplied values beginning '=' as formulas.
    value = str(text or '')
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value


def _num(value):
    s = str(value)
    return int(s) if s.isdigit() and not s.startswith('0') else _safe_text(s)


def _name_for_sheet(groups, index):
    start, end = groups[0].week, groups[-1].week
    label = f'{start}-{end}. Haftalar' if start != end else f'{start}. Hafta'
    return f'{index:02d} - {label}'[:31]


def build_workbook(students, department, course, instructor, title, weeks, signatures, midterm, final,
                   min_mark_mm=MARK_MIN_MM, font_size_pt=FONT_DEFAULT_PT,
                   font_family='Aptos', name_max_mm=NAME_MAX_DEFAULT_MM):
    """Return an A4 workbook + the exact dimensions/fonts for the PDF renderer."""
    min_mark_mm, font_size_pt, font_family, name_max_mm = validate_presentation(
        min_mark_mm, font_size_pt, font_family, name_max_mm)
    if not students:
        raise ValueError('Listede öğrenci yok')
    name_width = estimated_name_width_mm(students, font_size_pt, name_max_mm, font_family)
    fixed_widths = (SERIAL_WIDTH_MM, STUDENT_NUMBER_WIDTH_MM, name_width)
    segments = page_plan(weeks, signatures, midterm, final, min_mark_mm, fixed_widths)
    # Old size is now 2 pt smaller by default. All text groups follow the slider.
    size_offset = font_size_pt - 10
    def font_pt(previous, lower=6):
        return max(lower, previous + size_offset)
    wb = Workbook()
    wb.remove(wb.active)
    total_students = len(students)
    name_font_sizes = [fitted_name_font_pt(name, name_width, font_size_pt, font_family)
                       for _,name in students]
    physical_pages = math.ceil(total_students / STUDENTS_PER_PAGE)
    blank = physical_pages * STUDENTS_PER_PAGE - total_students
    last_row = HEADER_ROWS + total_students + blank
    layouts = []
    hairline = Side(style='thin', color=LIGHT_LINE)
    thickline = Side(style='medium', color=BURGUNDY)
    total_segments = len(segments)
    for i, groups in enumerate(segments, 1):
        ws = wb.create_sheet(_name_for_sheet(groups, i))
        width = mark_width_mm(groups, min_mark_mm, fixed_widths)
        total_marks = sum(g.span for g in groups)
        last_col = 3 + total_marks
        last_letter = get_column_letter(last_col)
        for col, mm in enumerate(fixed_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = mm_to_excel_chars(mm)
        for col in range(4, last_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = mm_to_excel_chars(width)
        for row, height in enumerate((27, 18, 24, 22, 23, 17), start=1):
            ws.row_dimensions[row].height = height
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        c = ws['A1']; c.value = title
        c.font = Font(name=font_family, size=font_pt(13 if len(title)<100 else 11.5, 8), bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=BURGUNDY)
        c.alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
        # Physical-width-based metadata layout. Two equal-ish blocks are used only
        # when each has room for its label; otherwise use two full-width rows.
        physical_columns = list(fixed_widths) + [width] * total_marks
        table_mm = sum(physical_columns)
        cutoffs = []
        for candidate_cut in range(1, last_col):
            left_mm = sum(physical_columns[:candidate_cut])
            right_mm = table_mm - left_mm
            if left_mm >= 55 and right_mm >= 55:
                cutoffs.append((abs(left_mm - right_mm), candidate_cut))
        compact_meta = table_mm < 160 or not cutoffs
        if compact_meta:
            ws.row_dimensions[2].height = 26
            ws.row_dimensions[3].height = 30
            for row, value in (
                (2, 'DERS KODU VE ADI: ' + str(course)),
                (3, 'ÖĞRETİM ELEMANI: ' + str(instructor)),
            ):
                ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last_col)
                c = ws.cell(row,1)
                c.value = _safe_text(value)
                c.font = Font(name=font_family,size=font_pt(8 if row==2 else 9),
                              bold=True,color=TEXT)
                c.alignment = Alignment(horizontal='center',vertical='center',
                                        wrap_text=True,shrink_to_fit=True)
                c.fill = PatternFill('solid',fgColor=PALE if row==2 else WHITE)
                c.border = Border(bottom=hairline,left=hairline,right=hairline)
        else:
            midpoint = min(cutoffs)[1]
            for row in (2,3):
                ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=midpoint)
                ws.merge_cells(start_row=row,start_column=midpoint+1,end_row=row,end_column=last_col)
            cells = (
                ('A2','DERS KODU VE ADI'),
                (f'{get_column_letter(midpoint+1)}2','ÖĞRETİM ELEMANI'),
                ('A3',course),
                (f'{get_column_letter(midpoint+1)}3',instructor),
            )
            for cell, value in cells:
                c=ws[cell];c.value=_safe_text(value)
                c.font=Font(name=font_family,size=font_pt(8 if c.row==2 else 10),
                            bold=True,color=TEXT)
                c.alignment=Alignment(horizontal='center',vertical='center',
                                      shrink_to_fit=True,wrap_text=True)
                c.fill=PatternFill('solid',fgColor=PALE if c.row==2 else WHITE)
                c.border=Border(bottom=hairline,left=hairline,right=hairline)
        ws.merge_cells(start_row=4,start_column=1,end_row=4,end_column=last_col)
        c=ws['A4'];c.value=_safe_text(department or 'BÖLÜM BİLGİSİ BELİRTİLMEDİ')
        c.font=Font(name=font_family,size=font_pt(11),bold=True,color=BURGUNDY)
        c.fill=PatternFill('solid',fgColor=GRAY)
        c.alignment=Alignment(horizontal='center',vertical='center',shrink_to_fit=True)
        for col, label in ((1,'SIRA'),(2,'ÖĞRENCİ NO'),(3,'AD SOYAD')):
            ws.merge_cells(start_row=5,start_column=col,end_row=6,end_column=col)
            c=ws.cell(5,col,label)
            c.font=Font(name=font_family,size=font_pt(9),bold=True,color=WHITE)
            c.fill=PatternFill('solid',fgColor=BURGUNDY)
            c.border=Border(left=hairline,right=hairline,bottom=hairline)
            c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        mark_col=4
        pdf_groups=[]
        exam_columns=[]
        for g in groups:
            start=mark_col
            end=mark_col+g.span-1
            if g.kind == 'week' and g.span == 1:
                # Exactly one weekly signature: merge the week heading vertically
                # across the two header rows; do not create a "1. İMZA" label.
                ws.merge_cells(start_row=5,start_column=start,end_row=6,end_column=start)
            elif g.span > 1:
                ws.merge_cells(start_row=5,start_column=start,end_row=5,end_column=end)
            c=ws.cell(5,start,g.label)
            c.font=Font(name=font_family,size=font_pt(8 if g.span>1 else 7),bold=True,color=WHITE if g.kind=='week' else TEXT)
            c.fill=PatternFill('solid',fgColor=BURGUNDY if g.kind=='week' else EXAM_GRAY)
            c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True,shrink_to_fit=True)
            c.border=Border(left=hairline,right=hairline,bottom=hairline)
            for mark_i in range(g.span):
                if g.kind == 'week' and g.span == 1:
                    continue
                x=ws.cell(6,start+mark_i)
                x.value=f'{mark_i+1}. İMZA' if g.kind=='week' else 'SINAV'
                x.font=Font(name=font_family,size=font_pt(7),bold=True,color=TEXT)
                x.alignment=Alignment(horizontal='center',vertical='center',shrink_to_fit=True)
                x.fill=PatternFill('solid',fgColor=PALE if g.kind=='week' else EXAM_GRAY)
                x.border=Border(left=hairline,right=hairline,bottom=hairline)
            if g.kind == 'exam':
                exam_columns.append(start)
            pdf_groups.append({'week':g.week,'kind':g.kind,'label':g.label,'span':g.span})
            mark_col=end+1
        for row in range(HEADER_ROWS+1,last_row+1):
            ws.row_dimensions[row].height=STUDENT_HEIGHT_PT
            idx=row-(HEADER_ROWS+1)
            student=students[idx] if idx<total_students else (None,None)
            serial=ws.cell(row,1,idx+1)
            number=ws.cell(row,2,_num(student[0]) if student[0] is not None else None)
            name=ws.cell(row,3,_safe_text(student[1]) if student[1] else None)
            for col in range(1,last_col+1):
                c=ws.cell(row,col)
                c.fill=PatternFill('solid',fgColor=(EXAM_GRAY if col in exam_columns else GRAY if idx%2 else WHITE))
                c.border=Border(left=hairline,right=hairline,bottom=hairline)
                c.alignment=Alignment(horizontal='left' if col==3 else 'center',vertical='center',wrap_text=col==3,shrink_to_fit=True)
                pt = name_font_sizes[idx] if col == 3 and idx < total_students else (font_size_pt if col == 3 else max(6,font_size_pt-1))
                c.font=Font(name='Aptos Narrow' if col==3 and font_family=='Aptos' else font_family,size=pt,bold=(col==3),color=TEXT)
        _print_setup(ws,last_col,last_row)
        pages=[]
        for page in range(physical_pages):
            page_rows=[]
            for j in range(STUDENTS_PER_PAGE):
                idx=page*STUDENTS_PER_PAGE+j
                num,name=students[idx] if idx<total_students else ('','')
                page_rows.append({'rank':idx+1, 'number':num, 'name':name,
                                  'nameFontPt':name_font_sizes[idx] if idx < total_students else font_size_pt})
            pages.append(page_rows)
        layouts.append({
            'name':ws.title,'title':title,'course':course,'instructor':instructor,
            'department':department,'groups':pdf_groups,'pages':pages,
            'capacity':STUDENTS_PER_PAGE,'markWidthMm':width,
            'fixedWidthsMm':list(fixed_widths),'rowHeightPt':STUDENT_HEIGHT_PT,
            'minMarkWidthMm':min_mark_mm,'nameMaxWidthMm':name_max_mm,
            'fontSizePt':font_size_pt,'fontFamily':font_family,
            'compactMeta':compact_meta,
            'printableWidthMm':PRINTABLE_MM,'totalWidthMm':round(sum(fixed_widths)+total_marks*width,3),
            'weekStart':groups[0].week,'weekEnd':groups[-1].week,
        })
    return wb,layouts,blank


def create_file(obs_path:Path,output_dir:Path,settings:dict):
    students,department,warnings = core.obs_oku(obs_path)
    weeks,signatures,midterm,final=validate_settings(settings['weeks'],settings['signatures'],settings['midterm'],settings['final'])
    title=(f'İSTANBUL RUMELİ ÜNİVERSİTESİ {settings["academicYear"]} EĞİTİM ÖĞRETİM YILI '
           f'{core.tr_buyuk(settings["semester"])} DÖNEMİ DERS YOKLAMA LİSTESİ')
    wb,layouts,blank=build_workbook(
        students,department,settings['course'],settings['instructor'],title,weeks,signatures,midterm,final,
        settings['minMarkWidthMm'],settings['fontSizePt'],settings['fontFamily'],settings['nameMaxWidthMm'])
    output_dir.mkdir(parents=True,exist_ok=True)
    course_code=re.sub(r'[^\w-]','',settings['course'].split('-',1)[0],flags=re.UNICODE)[:30] or 'DERS'
    branch=re.sub(r'[^\w-]','',obs_path.stem.rsplit('-',1)[-1],flags=re.UNICODE).upper()[:26] or 'SUBE'
    dest=output_dir / f'{course_code}-Yoklama-{branch}.xlsx'
    wb.save(dest)
    wb.close()
    return dest,students,department,warnings,layouts,blank
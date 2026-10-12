#!/usr/bin/env python3
"""OBS şube listelerinden ders yoklama listesi (Excel + PDF) üretir.

Haftalık imza sayısına göre şablon seçilir (Yoklama-1/2/3-Sign-Template...),
kopyalanır ve her sayfasına yalnızca şunlar yazılır: sıra numarası, öğrenci
numarası + ad soyad, bölüm başlığı (A4), ders ve öğretim elemanı. Öğrenci alanı,
son öğrencinin bulunduğu A4 sayfasının sonuna kadar uzatılır ya da kırpılır;
böylece geç kayıt olan öğrencilerin adlarını elle yazabileceği boş satırlar
kalır. Satırlar pijama (bir açık gri, bir beyaz) boyanır. Son olarak Excel ile
her sayfanın PDF'i alınır.

Kullanım:
    python yoklama_olustur.py 205-*.xlsx                      (3 imza; ders/hoca şablondaki)
    python yoklama_olustur.py 205-*.xlsx --imza 2 --ders "MAT205 - Lineer Cebir" --hoca "Dr. Öğr.Üyesi YASİN UCAKAN"
    python yoklama_olustur.py 205-*.xlsx --pdfsiz             (yalnızca Excel)

Çıktılar, OBS dosyalarının yanındaki "Yoklamalar" klasörüne yazılır:
    205-bb.xlsx  ->  Yoklamalar/MAT205-Yoklama-BB.xlsx
                     Yoklamalar/PDF/MAT205-Yoklama-BB - Vize Öncesi.pdf
                     Yoklamalar/PDF/MAT205-Yoklama-BB - Vize Sonrası.pdf
Tek sayfalı şablonda (1 imza) PDF adı Excel dosyasıyla aynıdır.
"""

import argparse
import glob
import re
import sys
from collections import Counter
from copy import copy
from pathlib import Path
from typing import NamedTuple

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import InvalidFileException
    from openpyxl.worksheet.pagebreak import Break, RowBreak
except ImportError:
    sys.exit("openpyxl gerekli. Kurmak için: python -m pip install openpyxl")

SABLON_DESENI = "Yoklama-{imza}-Sign-Template*.xlsx"
VARSAYILAN_IMZA = 3
CIKTI_KLASORU = "Yoklamalar"
PDF_KLASORU = "PDF"

# Üç şablonda da ortak olanlar; geri kalan düzen şablondan okunur (sablon_duzeni).
BOLUM_HUCRESI = "A4"
SIRA_SUTUNU, NO_SUTUNU, AD_SUTUNU = 1, 2, 3
DERS_BASLIGI = "Ders Kodu"        # 2. satırdaki başlık; değer hemen altına (3. satır) yazılır
HOCA_BASLIGI = "Öğretim Elemanı"

# OBS başlıkları "Öğrenci No_9898YU" biçiminde gelir; sondaki ek yok sayılır.
OBS_SUTUNLARI = {"no": "Öğrenci No", "ad": "Adı", "soyad": "Soyadı", "program": "Program"}

KALIN = Side(style="medium")
# Pijama rengi: gri sınav bloklarından (%15) belirgin açık, siyah-beyaz baskıda da seçilir.
PIJAMA = PatternFill("solid", fgColor="EDEDED")
DOLGUSUZ = PatternFill(fill_type=None)

# İsim puntosu şablondan gelir (kalın 16). İsim o puntoda 2 satıra sığmıyorsa (satır yüksekliği 2 satır
# alıyor) yalnızca o hücrenin puntosu küçültülür. Sığıp sığmadığı, Excel'de ölçülmüş harf genişlikleriyle
# tahmin edilir (kalın Aptos Narrow, 16 punto, sütun genişliği biriminde).
HARF_GENISLIGI = {" ": 0.57, "'": 0.64, ".": 0.87, "I": 0.87, "İ": 0.87, "-": 0.96, "J": 1.02,
                  **dict.fromkeys("FLTZ", 1.46), **dict.fromkeys("EPSŞXY", 1.63), "V": 1.76,
                  **dict.fromkeys("ABKR", 1.79), **dict.fromkeys("CÇDGĞUÜ", 1.96), **dict.fromkeys("HN", 2.06),
                  **dict.fromkeys("OÖQ", 2.13), "M": 2.29, "W": 2.64}
VARSAYILAN_HARF = 1.8
EN_KUCUK_PUNTO = 10


class YoklamaHatasi(Exception):
    pass


class Duzen(NamedTuple):
    ilk: int           # ilk öğrenci satırı (yazdırma başlık satırlarından hemen sonra)
    son: int           # şablondaki son öğrenci satırı
    sayfa_basina: int  # bir A4'e sığan öğrenci satırı (şablondaki ilk sayfa sonundan)
    son_sutun: int     # tablonun son sütunu
    bloklar: dict      # gri VİZE / FİNAL blokları: (ilk sütun, son sütun) -> üst satır


def tr_buyuk(metin):
    """Türkçe büyük harf: str.upper() 'i' harfini 'I' yapar, burada 'İ' olur."""
    return metin.replace("i", "İ").replace("ı", "I").upper()


def temizle(deger):
    return " ".join(str(deger or "").split())


def obs_oku(yol):
    """OBS listesinden (numara, ad soyad) çiftlerini, bölüm adını ve uyarıları döndürür."""
    satirlar = load_workbook(yol, data_only=True).active.iter_rows(values_only=True)
    basliklar = [re.sub(r"_[^_]*$", "", temizle(b)).casefold() for b in next(satirlar, ())]
    try:
        sutun = {k: basliklar.index(ad.casefold()) for k, ad in OBS_SUTUNLARI.items()}
    except ValueError:
        raise YoklamaHatasi("OBS öğrenci listesi değil (Öğrenci No / Adı / Soyadı / Program sütunları yok), atlandı")

    ogrenciler, programlar = [], []
    for satir in satirlar:
        no = temizle(satir[sutun["no"]])
        if not no:
            continue
        ad_soyad = tr_buyuk(temizle(f"{satir[sutun['ad']] or ''} {satir[sutun['soyad']] or ''}"))
        if not ad_soyad:
            raise YoklamaHatasi(f"{no} numaralı öğrencinin adı soyadı boş")
        ogrenciler.append((no, ad_soyad))
        program = temizle(satir[sutun["program"]])
        if program and program not in programlar:
            programlar.append(program)

    if not ogrenciler:
        raise YoklamaHatasi("listede hiç öğrenci yok")
    uyarilar = [f"{no} numarası listede {adet} kez geçiyor"
                for no, adet in Counter(no for no, _ in ogrenciler).items() if adet > 1]
    if len(programlar) > 1:
        uyarilar.append("listede birden fazla program var, başlığa hepsi yazıldı")
    return ogrenciler, " / ".join(tr_buyuk(p) for p in programlar), uyarilar


def sablon_bul(imza):
    adaylar = sorted(Path(__file__).parent.glob(SABLON_DESENI.format(imza=imza)))
    if not adaylar:
        raise YoklamaHatasi(f"{imza} imzalı şablon bulunamadı ({SABLON_DESENI.format(imza=imza)})")
    return adaylar[0]


def bilgi_hucresi(ws, baslik):
    """2. satırdaki başlığın (ör. "Ders Kodu ve İsmi") altındaki hücre."""
    for hucre in ws[2]:
        if baslik in str(hucre.value or ""):
            return ws.cell(3, hucre.column)
    raise YoklamaHatasi(f"şablonun '{ws.title}' sayfasında '{baslik}' başlığı bulunamadı")


def sablon_duzeni(ws):
    if not ws.print_title_rows:
        raise YoklamaHatasi(f"şablonun '{ws.title}' sayfasında yazdırma başlık satırları tanımlı değil")
    ilk = int(ws.print_title_rows.split(":")[-1].strip("$")) + 1
    # Öğrenci satırlarına uzanan birleşik hücreler gri VİZE / FİNAL bloklarıdır.
    uzananlar = [r for r in ws.merged_cells.ranges if r.max_row >= ilk]
    if not uzananlar:
        raise YoklamaHatasi(f"şablonun '{ws.title}' sayfasında sınav bloğu bulunamadı")
    bloklar = {}
    for r in uzananlar:
        bloklar[(r.min_col, r.max_col)] = min(r.min_row, bloklar.get((r.min_col, r.max_col), r.min_row))
    sayfa_sonlari = sorted(b.id for b in ws.row_breaks.brk if b.id >= ilk)
    if not sayfa_sonlari:
        raise YoklamaHatasi(f"şablonun '{ws.title}' sayfasında sayfa sonu tanımlı değil; "
                            "bir A4'e kaç öğrenci sığdığı buradan okunuyor")
    baslik = next((r for r in ws.merged_cells.ranges if r.min_row == r.min_col == 1), None)
    return Duzen(ilk=ilk, son=max(r.max_row for r in uzananlar), sayfa_basina=sayfa_sonlari[0] - ilk + 1,
                 son_sutun=baslik.max_col if baslik else ws.max_column, bloklar=bloklar)


def gri_blok(ws, r1, c1, r2, c2, stil, deger, kenar):
    hucre = ws.cell(r1, c1)
    hucre._style = copy(stil)
    hucre.border = kenar
    hucre.value = deger
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def sayfa_duzeni(ws, d, yeni_son):
    """Öğrenci alanını yeni_son satırında bitirir, gri sınav bloklarını her A4 sayfası için ayrı birleştirir.

    Tek parça blokta "VİZE SINAVI" yazısı bloğun ortasına, yani çoğu zaman
    2. sayfaya düşüyordu; sayfa başına ayrı blok olunca her sayfada görünür.
    Bloğun başlık satırlarındaki parçası ayrı tutulur ve aradaki çizgi kaldırılır;
    böylece her sayfada tekrarlanan başlıkla birlikte kesintisiz bir sütun olur.
    Satır silinip eklendiği için pijama boyaması da burada baştan yapılır.
    """
    bloklar = [(c1, c2, ust, ws.cell(ust, c1).value, copy(ws.cell(ust, c1)._style))
               for (c1, c2), ust in d.bloklar.items()]
    for r in [r for r in ws.merged_cells.ranges if r.max_row >= d.ilk]:
        ws.unmerge_cells(r.coord)

    if yeni_son < d.son:
        # Aradaki satırlar silinince alt kenarı kalın olan son satır yukarı kayar.
        ws.delete_rows(yeni_son, d.son - yeni_son)
    elif yeni_son > d.son:
        # Son satırın önüne, bir önceki satırın biçiminde satırlar eklenir; son satır aşağı kayar.
        ws.insert_rows(d.son, yeni_son - d.son)
        for satir in range(d.son, yeni_son):
            for sutun in range(1, d.son_sutun + 1):
                ws.cell(satir, sutun)._style = copy(ws.cell(d.son - 1, sutun)._style)
    yukseklik = ws.row_dimensions[d.ilk].height
    for satir in range(d.ilk, max(yeni_son, d.son) + 1):
        if satir <= yeni_son:
            ws.row_dimensions[satir].height = yukseklik
        else:
            ws.row_dimensions.pop(satir, None)
    # Excel'in "kalın kenarlı satır" işaretleri eski düzenden kalıyor; başlığın son satırında kalırsa
    # Excel sonraki sayfalarda başlıkla ilk satır arasında ince bir boşluk bırakıyor.
    for boyut in ws.row_dimensions.values():
        boyut.thickTop = boyut.thickBot = False

    ws.row_breaks = RowBreak()
    sayfa_sonlari = list(range(d.ilk - 1 + d.sayfa_basina, yeni_son, d.sayfa_basina)) + [yeni_son]
    for i, sayfa_sonu in enumerate(sayfa_sonlari):
        for c1, c2, ust, etiket, stil in bloklar:
            if i == 0 and ust < d.ilk:
                # Üst kenar gerekmez: üstteki bölüm şeridinin kalın alt kenarı zaten çiziyor.
                gri_blok(ws, ust, c1, d.ilk - 1, c2, stil, None, Border(left=KALIN, right=KALIN))
                # Boş başlık parçası dikey yazı ayarı taşımasın: Excel onu sonraki sayfalarda bozuk basabiliyor.
                ws.cell(ust, c1).alignment = Alignment()
            bas = d.ilk if i == 0 else sayfa_sonlari[i - 1] + 1
            gri_blok(ws, bas, c1, sayfa_sonu, c2, stil, etiket, Border(left=KALIN, right=KALIN, bottom=KALIN))
        if sayfa_sonu < yeni_son:
            ws.row_breaks.append(Break(id=sayfa_sonu))
    ws.print_area = f"A1:{get_column_letter(d.son_sutun)}{yeni_son}"

    sinav_sutunlari = {s for c1, c2, *_ in bloklar for s in range(c1, c2 + 1)}
    for satir in range(d.ilk, yeni_son + 1):
        dolgu = PIJAMA if (satir - d.ilk) % 2 else DOLGUSUZ
        for sutun in range(1, d.son_sutun + 1):
            if sutun not in sinav_sutunlari:
                ws.cell(satir, sutun).fill = dolgu
    return len(sayfa_sonlari)


def satir_sayisi(ad, punto, kapasite):
    """İsmin sütunda kaç satır tutacağı; Excel gibi kelime kelime alt satıra geçirir."""
    olcek = punto / 16
    bosluk = HARF_GENISLIGI[" "] * olcek
    satirlar, dolu = 1, 0.0
    for kelime in ad.split():
        genislik = sum(HARF_GENISLIGI.get(h, VARSAYILAN_HARF) for h in kelime) * olcek
        if dolu and dolu + bosluk + genislik > kapasite:
            satirlar, dolu = satirlar + 1, genislik
        else:
            dolu += (bosluk if dolu else 0) + genislik
    return satirlar


def isim_puntosu(ad, punto, kapasite):
    while punto > EN_KUCUK_PUNTO and satir_sayisi(ad, punto, kapasite) > 2:
        punto -= 1
    return punto


def numara(no):
    # Excel'in "metin olarak saklanan sayı" uyarısı çıkmasın diye sayıya çevrilir.
    return int(no) if no.isdigit() and not no.startswith("0") else no


def yoklama_olustur(obs_yolu, sablon_yolu, ders, hoca):
    ogrenciler, bolum, uyarilar = obs_oku(obs_yolu)
    wb = load_workbook(sablon_yolu)
    for ws in wb.worksheets:
        d = sablon_duzeni(ws)
        yeni_son = d.ilk - 1 + -(-len(ogrenciler) // d.sayfa_basina) * d.sayfa_basina
        sayfa = sayfa_duzeni(ws, d, yeni_son)
        # Excel'deki sütun genişliği ~0.7 birim kenar payı içerir; metne kalan alan biraz daha dar.
        kapasite = ws.column_dimensions[get_column_letter(AD_SUTUNU)].width - 1.7

        for i, satir in enumerate(range(d.ilk, yeni_son + 1)):
            no, ad_soyad = ogrenciler[i] if i < len(ogrenciler) else (None, None)
            # Sıra numarası boş satırlara da yazılır: geç kayıt olan öğrenci sıradakine yazar.
            ws.cell(satir, SIRA_SUTUNU).value = i + 1
            ws.cell(satir, NO_SUTUNU).value = numara(no) if no else None
            hucre = ws.cell(satir, AD_SUTUNU)
            hucre.value = ad_soyad
            if ad_soyad:
                punto = isim_puntosu(ad_soyad, hucre.font.sz, kapasite)
                if punto != hucre.font.sz:
                    font = copy(hucre.font)
                    font.sz = punto
                    hucre.font = font
        if bolum:
            ws[BOLUM_HUCRESI] = bolum
        bilgi_hucresi(ws, DERS_BASLIGI).value = ders
        bilgi_hucresi(ws, HOCA_BASLIGI).value = hoca

    ders_kodu = re.sub(r'[\s<>:"/\\|?*]', "", ders.split("-")[0])
    sube = obs_yolu.stem.rsplit("-", 1)[-1].upper()
    cikti = obs_yolu.parent / CIKTI_KLASORU / f"{ders_kodu or 'DERS'}-Yoklama-{sube}.xlsx"
    cikti.parent.mkdir(exist_ok=True)
    wb.save(cikti)
    bos_satir = yeni_son - d.ilk + 1 - len(ogrenciler)
    return cikti, len(ogrenciler), bos_satir, bolum, sayfa, uyarilar


def pdf_al(xlsx_dosyalari):
    """Excel'in kendi PDF çıktısıyla (Yazdır > PDF ile aynı sonuç) her sayfayı ayrı PDF yapar.

    (pdf yolu, hata mesajı ya da None) çiftleri üretir.
    """
    try:
        import pywintypes
        import win32com.client
    except ImportError:
        raise YoklamaHatasi("PDF için pywin32 gerekli: python -m pip install pywin32")
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
    except pywintypes.com_error:
        raise YoklamaHatasi("Excel başlatılamadı, PDF alınamadı")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        for xlsx in xlsx_dosyalari:
            xlsx = xlsx.resolve()  # Excel göreli yolu kendi klasörüne göre çözer
            klasor = xlsx.parent / PDF_KLASORU
            klasor.mkdir(exist_ok=True)
            wb = excel.Workbooks.Open(str(xlsx), 0, True)
            try:
                sayfalar = list(wb.Worksheets)
                for ws in sayfalar:
                    pdf = klasor / (f"{xlsx.stem}.pdf" if len(sayfalar) == 1 else f"{xlsx.stem} - {ws.Name}.pdf")
                    try:
                        ws.ExportAsFixedFormat(0, str(pdf))
                    except pywintypes.com_error as e:
                        aciklama = (e.excepinfo or [None] * 3)[2] or e.strerror
                        yield pdf, f"yazılamadı ({aciklama}); PDF başka bir programda açıksa kapatın"
                    else:
                        yield pdf, None
            finally:
                wb.Close(False)
    finally:
        excel.Quit()


def sor(soru, varsayilan="", secenekler=None):
    ek = f" [{varsayilan}]" if varsayilan else ""
    while True:
        try:
            cevap = input(f"{soru}{ek}: ").strip() or varsayilan
        except EOFError:
            return varsayilan
        if cevap and (not secenekler or cevap in secenekler):
            return cevap


def dosyalari_topla(desenler):
    # PowerShell ve cmd "205-*.xlsx" gibi desenleri açmadığı için burada açılır.
    dosyalar = []
    for desen in desenler:
        eslesenler = (sorted(glob.glob(desen)) if any(c in desen for c in "*?[") else []) or [desen]
        dosyalar += [Path(d) for d in eslesenler if not Path(d).name.startswith("~$")]
    return dosyalar


def main():
    for akis in (sys.stdout, sys.stderr):
        akis.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="OBS şube listelerinden ders yoklama listesi (Excel + PDF) üretir.")
    parser.add_argument("obs_dosyalari", nargs="*", help="OBS'den indirilen şube listeleri (ör. 205-bb.xlsx)")
    parser.add_argument("--imza", type=int, choices=(1, 2, 3), default=VARSAYILAN_IMZA,
                        help=f"haftalık imza (ders saati) sayısı; şablonu belirler (varsayılan: {VARSAYILAN_IMZA})")
    parser.add_argument("--sablon", type=Path, help="şablonu elle seçmek için (verilirse --imza yok sayılır)")
    parser.add_argument("--ders", help='ör. "MAT205 - Lineer Cebir" (varsayılan: şablondaki)')
    parser.add_argument("--hoca", help='ör. "Dr. Öğr.Üyesi YASİN UCAKAN" (varsayılan: şablondaki)')
    parser.add_argument("--pdfsiz", action="store_true", help="PDF alma, yalnızca Excel dosyalarını üret")
    parser.add_argument("--sor", action="store_true", help="imza sayısı, ders ve öğretim elemanını ekrandan sor")
    args = parser.parse_args()

    dosyalar = dosyalari_topla(args.obs_dosyalari)
    if not dosyalar:
        parser.print_help()
        return 1
    try:
        if args.sor:
            print("Parantez içindekini kullanmak için Enter'a basın.")
            args.imza = int(sor("Haftalık imza sayısı (1/2/3)", str(args.imza), {"1", "2", "3"}))
        sablon = args.sablon or sablon_bul(args.imza)
        if not sablon.is_file():
            raise YoklamaHatasi(f"şablon bulunamadı: {sablon}")
        ilk_sayfa = load_workbook(sablon).worksheets[0]
        ders = args.ders or temizle(bilgi_hucresi(ilk_sayfa, DERS_BASLIGI).value)
        hoca = args.hoca or temizle(bilgi_hucresi(ilk_sayfa, HOCA_BASLIGI).value)
        if args.sor:
            ders = sor("Ders kodu ve adı", ders)
            hoca = sor("Öğretim elemanı", hoca)
            print()
        if not ders or not hoca:
            raise YoklamaHatasi("şablonda ders / öğretim elemanı boş; --ders ve --hoca ile verin")
    except YoklamaHatasi as e:
        print(f"HATA  {e}")
        return 1
    print(f"Şablon: {sablon.name}")

    uretilenler, hatali = [], 0
    for yol in dosyalar:
        try:
            if not yol.is_file():
                raise YoklamaHatasi("dosya bulunamadı")
            cikti, adet, bos, bolum, sayfa, uyarilar = yoklama_olustur(yol, sablon, ders, hoca)
        except YoklamaHatasi as e:
            hata = str(e)
        except InvalidFileException:
            hata = "okunamadı (.xlsx dosyası olmalı)"
        except PermissionError:
            hata = "çıktı dosyası Excel'de açık olabilir; kapatıp tekrar deneyin"
        else:
            print(f"OK    {yol.name} -> {CIKTI_KLASORU}/{cikti.name}  "
                  f"({adet} öğrenci + {bos} boş satır, {bolum}, her sayfa {sayfa} A4)")
            for uyari in uyarilar:
                print(f"      uyarı: {uyari}")
            uretilenler.append(cikti)
            continue
        print(f"HATA  {yol.name}: {hata}")
        hatali += 1

    if uretilenler and not args.pdfsiz:
        try:
            for pdf, hata in pdf_al(uretilenler):
                if hata:
                    print(f"HATA  {pdf.name}: {hata}")
                    hatali += 1
                else:
                    print(f"PDF   {CIKTI_KLASORU}/{PDF_KLASORU}/{pdf.name}")
        except YoklamaHatasi as e:
            print(f"HATA  {e}")
            hatali += 1
    return 1 if hatali else 0


if __name__ == "__main__":
    sys.exit(main())
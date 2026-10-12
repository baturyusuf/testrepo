"""Combinatorial layout and generated-workbook integration tests."""
import math
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import sys

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import adaptive_layout as layout
import web_bridge


def test_all_combinations():
    checked=0
    for weeks in range(10,21):
        for marks in range(1,6):
            for vize,final in [(8,min(16,weeks)), (1,weeks), (weeks-1,weeks), (2,min(9,weeks))]:
                pages=layout.page_plan(weeks,marks,vize,final)
                expanded=[g for page in pages for g in page]
                assert [g.week for g in expanded] == list(range(1,weeks+1))
                assert len([g for g in expanded if g.kind=='exam']) == 2
                assert all(len(page)>0 for page in pages)
                for page in pages:
                    width=layout.mark_width_mm(page)
                    assert 8.8 <= width <= 18
                    assert sum(g.span for g in page)*width <= 195.01
                checked+=1
    assert len(layout.page_plan(20,5,8,16)) == 5
    assert len(layout.page_plan(16,3,8,16)) == 2
    assert len(layout.page_plan(16,1,8,16)) == 1
    return checked


def test_workbook_and_bridge():
    checked=0
    with TemporaryDirectory() as t:
        work=Path(t)
        (work/'uploads'/'1').mkdir(parents=True)
        (work/'uploads'/'2').mkdir(parents=True)
        for n,count in [(1,1),(2,105)]:
            wb=Workbook(); ws=wb.active
            ws.append(['Öğrenci No_42','Adı_42','Soyadı_42','Program_42'])
            for k in range(count): ws.append((f'0{k+1:07d}','AYŞE','YILMAZ','BİLGİSAYAR MÜHENDİSLİĞİ'))
            wb.save(work/'uploads'/str(n)/f'205-b{n}.xlsx')
        web_bridge.WORK=work
        for weeks,signatures,vize,final in [(16,1,8,16),(16,3,8,16),(10,2,4,9),(20,5,8,16),(20,1,6,19)]:
            data={
                'weeks':weeks, 'signatures':signatures, 'midterm':vize,'final':final,
                'course':'BLG103 - Programlama I','instructor':'Dr. Öğr. Üyesi Deneme',
                'academicYear':'2026-2027','semester':'Güz',
                'files':[{'name':'205-b1.xlsx'},{'name':'205-b2.xlsx'}]
            }
            results=json.loads(web_bridge.generate(json.dumps(data)))
            assert not results['errors'],results['errors']
            assert len(results['results'])==2
            for original,output in zip((1,105),results['results']):
                assert output['count']==original
                assert output['blankRows']==math.ceil(original/17)*17-original
                wb=load_workbook(output['file'])
                assert len(wb.worksheets)==len(layout.page_plan(weeks,signatures,vize,final))
                assert len(output['sheets'])==len(wb.worksheets)
                all_groups=[g for sh in output['sheets'] for g in sh['groups']]
                assert [g['week'] for g in all_groups]==list(range(1,weeks+1))
                assert sum(g['kind']=='exam' for g in all_groups)==2
                for ws,sh in zip(wb.worksheets,output['sheets']):
                    assert '2026-2027' in ws['A1'].value
                    assert ws.page_setup.orientation == 'landscape'
                    assert ws.page_setup.fitToWidth == 1
                    assert ws.print_title_rows == '$1:$6'
                    assert len(sh['pages'])==math.ceil(original/17)
                    assert len(ws.row_breaks.brk)==max(0,len(sh['pages'])-1)
                    assert sh['markWidthMm'] >= 8.8
                    assert sh['totalWidthMm'] <= 289.001
                    assert ws['C7'].value=='AYŞE YILMAZ'
                    assert ws['B7'].value == '00000001'
                    assert ws.row_dimensions[7].height>=layout.MIN_STUDENT_HEIGHT_PT
                    assert ws.row_dimensions[7].height<=layout.MAX_STUDENT_HEIGHT_PT
                wb.close()
                checked+=1
        try:
            web_bridge.generate(json.dumps({**data,'midterm':20,'final':8}))
        except ValueError: pass
        else: raise AssertionError('Invalid exams accepted')
    return checked

if __name__=='__main__':
    print('Cases checked:',test_all_combinations(),'layout combinations,',test_workbook_and_bridge(),'workbooks.')